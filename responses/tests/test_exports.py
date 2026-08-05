import csv
import io
import json

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from sharing.models import SurveyCollaborator
from surveys.models import Question, Survey
from surveys.publication import publish_survey

from responses.models import Submission


class ResponseExportTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.owner = User.objects.create_user(email='export-owner@example.com')
        self.respondent = User.objects.create_user(email='private-export-user@example.com')
        self.viewer = User.objects.create_user(email='export-viewer@example.com')
        self.outsider = User.objects.create_user(email='export-outsider@example.com')
        self.survey = Survey.objects.create(
            owner=self.owner,
            title='Unicode export study',
            summary='Filtered export fixtures.',
        )
        draft = self.survey.draft_version
        Question.objects.create(
            section=draft.sections.get(),
            type=Question.Type.LONG_TEXT,
            prompt='Describe your experience',
        )
        self.version, _ = publish_survey(self.survey.id, self.owner, draft.revision)
        self.survey.refresh_from_db()
        self.question = self.version.sections.get().questions.get()
        SurveyCollaborator.objects.create(
            survey=self.survey,
            user=self.viewer,
            role=SurveyCollaborator.Role.VIEWER,
            added_by=self.owner,
        )

    def complete_response(self):
        browser = self.client_class()
        browser.force_login(self.respondent)
        browser.post(reverse('respond_survey', args=[self.survey.slug]))
        submission = Submission.objects.get()
        browser.post(
            reverse('response_form', args=[submission.id]),
            {f'q_{self.question.id}': '=1+1\nবাংলা response'},
        )
        submission.refresh_from_db()
        return submission

    def body(self, response):
        return b''.join(response.streaming_content)

    def test_csv_exports_filtered_unicode_multiline_answers(self):
        completed = self.complete_response()
        Submission.objects.create(
            survey=self.survey,
            version=self.version,
            session_key_hash='f' * 64,
            presentation={'sections': []},
        )
        self.client.force_login(self.owner)

        response = self.client.get(
            reverse('response_export_csv', args=[self.survey.id]),
            {f'answer_{self.question.id}': 'response'},
        )

        rows = list(
            csv.reader(io.StringIO(self.body(response).decode('utf-8-sig')))
        )
        self.assertEqual(response.status_code, 200)
        # A single clean table: one header row, one data row. Version metadata
        # lives in the JSON/Excel exports, not as extra CSV rows.
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0][0], 'Submitted at')
        self.assertIn('Describe your experience', rows[0][1])
        self.assertEqual(len(rows[1]), 2)
        self.assertEqual(rows[1][0], completed.completed_at.isoformat())
        self.assertIn("'=1+1\nবাংলা response", rows[1])
        self.assertNotIn(self.respondent.email, rows[1])

    def test_json_preserves_structured_records_without_platform_identity(self):
        submission = self.complete_response()
        self.client.force_login(self.owner)

        response = self.client.get(reverse('response_export_json', args=[self.survey.id]))
        payload = json.loads(self.body(response))

        self.assertEqual(payload['survey']['id'], str(self.survey.id))
        self.assertEqual(
            payload['responses'][0]['submitted_at'],
            submission.completed_at.isoformat(),
        )
        self.assertEqual(
            payload['responses'][0]['answers'][str(self.question.id)],
            '=1+1\nবাংলা response',
        )
        self.assertNotIn('identity_email', payload['responses'][0])
        self.assertNotIn('response_id', payload['responses'][0])
        self.assertNotIn(self.respondent.email, self.body(response).decode())

    def test_viewer_can_export_but_outsider_cannot(self):
        self.complete_response()
        url = reverse('response_export_json', args=[self.survey.id])
        self.client.force_login(self.viewer)
        self.assertEqual(self.client.get(url).status_code, 200)
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(url).status_code, 404)

    def test_excel_exports_filtered_rows_as_safe_unicode_cells(self):
        submission = self.complete_response()
        Submission.objects.create(
            survey=self.survey,
            version=self.version,
            session_key_hash='e' * 64,
            presentation={'sections': []},
        )
        self.client.force_login(self.owner)

        response = self.client.get(
            reverse('response_export_excel', args=[self.survey.id]),
            {f'answer_{self.question.id}': 'response'},
        )
        workbook = load_workbook(io.BytesIO(self.body(response)), read_only=True)
        rows = list(workbook['Responses'].iter_rows(values_only=False))
        values = [cell.value for cell in rows[1]]

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0][0].value, 'Submitted at')
        self.assertIn('Describe your experience', rows[0][1].value)
        self.assertEqual(len(values), 2)
        self.assertEqual(values[0], submission.completed_at.isoformat())
        self.assertIn("'=1+1\nবাংলা response", values)
        formula_safe_cell = next(cell for cell in rows[1] if cell.value == "'=1+1\nবাংলা response")
        self.assertEqual(formula_safe_cell.data_type, 's')
        self.assertNotIn(self.respondent.email, values)
        self.assertEqual(workbook['About']['B1'].value, self.survey.title)
        self.assertEqual(workbook['About']['B2'].value, self.survey.title)
        self.assertEqual(workbook['About']['B3'].value, 1)

    def test_invalid_filters_return_bad_request(self):
        self.client.force_login(self.owner)

        response = self.client.get(
            reverse('response_export_csv', args=[self.survey.id]),
            {'submitted_from': '2026-02-02', 'submitted_to': '2026-01-01'},
        )

        self.assertEqual(response.status_code, 400)
