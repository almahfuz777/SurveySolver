import csv
from tempfile import SpooledTemporaryFile

from django.core.serializers.json import DjangoJSONEncoder
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill

from surveys.models import Question, SurveyVersion
from surveys.presentation import resolved_version

from .dashboard import filter_response_sheet, format_answer_cell
from .forms import ResponseSheetFilterForm
from .models import Submission


TIMESTAMP_COLUMN = 'Submitted at'


class Echo:
    def write(self, value):
        return value


def filtered_export_data(query_parameters, survey):
    versions = survey.versions.exclude(status=SurveyVersion.Status.DRAFT)
    version_id = query_parameters.get('sheet_version')
    if version_id:
        version = versions.filter(pk=version_id).first()
    else:
        version = versions.filter(status=SurveyVersion.Status.PUBLISHED).first()
        if version is None:
            version = versions.order_by('-number').first()
    if version is None:
        return ResponseSheetFilterForm(query_parameters, questions=[]), None, None, None
    questions = response_sheet_questions(survey, version_id=version.id)
    form = ResponseSheetFilterForm(
        query_parameters,
        questions=questions,
    )
    if not form.is_valid():
        return form, None, None, version
    submissions = filter_response_sheet(
        survey.submissions.filter(
            version=version,
            status=Submission.Status.COMPLETED,
        ).select_related('version').prefetch_related(
            'answers__question__section',
        ),
        form.cleaned_data,
        questions,
    )
    return form, submissions, questions, version


def response_sheet_questions(survey, *, version_id=None):
    version = (
        survey.versions.exclude(status=SurveyVersion.Status.DRAFT)
        .filter(pk=version_id)
        .first()
        if version_id
        else survey.versions.filter(status=SurveyVersion.Status.PUBLISHED).first()
    )
    if version is None:
        return []
    questions = [
        question
        for section in resolved_version(version)
        for question in section.questions.all()
    ]
    for question in questions:
        question.sheet_label = question_column(question)
    return questions


def _csv_safe(value):
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@', '\t', '\r')):
        return f"'{value}"
    return value


def _excel_safe(value):
    value = _csv_safe(value)
    if not isinstance(value, str):
        return value
    value = ILLEGAL_CHARACTERS_RE.sub('', value)
    if len(value) > 32767:
        value = f'{value[:32740]}… [truncated for Excel]'
    return value


def _blank_if_none(value):
    return '' if value is None else value


def response_sheet_record(submission, questions):
    answer_by_question = {
        answer.question_id: format_answer_cell(answer)
        for answer in submission.answers.all()
    }
    return {
        'submitted_at': submission.completed_at.isoformat(),
        'answers': {
            str(question.id): answer_by_question.get(question.id)
            for question in questions
        },
    }


def question_column(question):
    return question.prompt


def iter_csv(submissions, questions):
    # Keep the CSV a single clean table (header + rows) so it imports directly
    # into spreadsheets and pandas. Version metadata lives in the JSON/Excel
    # exports, which have somewhere structured to put it.
    writer = csv.writer(Echo())
    yield '\ufeff' + writer.writerow(
        (TIMESTAMP_COLUMN, *(_csv_safe(question.sheet_label) for question in questions))
    )
    for submission in submissions.iterator(chunk_size=500):
        record = response_sheet_record(submission, questions)
        yield writer.writerow(
            [_csv_safe(record['submitted_at'])]
            + [
                _csv_safe(
                    _blank_if_none(record['answers'].get(str(question.id)))
                )
                for question in questions
            ]
        )


def iter_json(survey, submissions, questions, version):
    encoder = DjangoJSONEncoder(ensure_ascii=False, separators=(',', ':'))
    yield '{"survey":'
    yield encoder.encode(
        {
            'id': str(survey.id),
            'current_title': survey.title,
            'version': version.number,
            'title_snapshot': version.title_snapshot or survey.title,
        }
    )
    yield ',"questions":'
    yield encoder.encode(
        [
            {'id': str(question.id), 'label': question.sheet_label}
            for question in questions
        ]
    )
    yield ',"generated_at":'
    yield encoder.encode(timezone.now().isoformat())
    yield ',"responses":['
    first = True
    for submission in submissions.iterator(chunk_size=500):
        if not first:
            yield ','
        yield encoder.encode(response_sheet_record(submission, questions))
        first = False
    yield ']}'


def build_excel(survey, submissions, questions, version):
    output = SpooledTemporaryFile(max_size=5 * 1024 * 1024, mode='w+b')
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet('Responses')
    worksheet.freeze_panes = 'A2'
    headers = (TIMESTAMP_COLUMN, *(question.sheet_label for question in questions))
    header_cells = []
    for header in headers:
        cell = WriteOnlyCell(worksheet, value=_excel_safe(header))
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='312E81')
        header_cells.append(cell)
    worksheet.append(header_cells)
    for submission in submissions.iterator(chunk_size=500):
        record = response_sheet_record(submission, questions)
        worksheet.append(
            [_excel_safe(record['submitted_at'])]
            + [
                _excel_safe(
                    _blank_if_none(record['answers'].get(str(question.id)))
                )
                for question in questions
            ]
        )
    metadata = workbook.create_sheet('About')
    metadata.append(('Current survey title', _excel_safe(survey.title)))
    metadata.append(
        ('Publication title', _excel_safe(version.title_snapshot or survey.title))
    )
    metadata.append(('Version', version.number))
    metadata.append(('Survey ID', str(survey.id)))
    metadata.append(('Generated at', timezone.now().isoformat()))
    workbook.save(output)
    output.seek(0)
    return output
